from io import BytesIO

from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.authtoken.models import Token
from rest_framework.test import APITestCase

from catalog.models import Department, Room, Section, Subject
from timetable.models import Assignment, ScheduledClass, ScheduleRun
from users.models import Professors

User = get_user_model()


class ExportTestCase(APITestCase):
    def setUp(self):
        # Registrar
        self.registrar = User.objects.create_user(
            username="reg", password="pass12345", is_staff=True
        )
        self.reg_token = Token.objects.create(user=self.registrar)

        # Professor
        self.dept = Department.objects.create(name="CS")
        self.prof_user = User.objects.create_user(
            username="prof1", password="pass12345"
        )
        self.prof = Professors.objects.create(
            user=self.prof_user, department=self.dept
        )
        self.prof_token = Token.objects.create(user=self.prof_user)

        # Schedule data
        self.subject = Subject.objects.create(
            code="CS101", title="Intro", units=3
        )
        self.section = Section.objects.create(name="BSIT-3A", headcount=30)
        self.room = Room.objects.create(name="R101", capacity=40)
        self.assignment = Assignment.objects.create(
            prof=self.prof, subject=self.subject, section=self.section
        )
        self.run = ScheduleRun.objects.create(
            algorithm="greedy", status=ScheduleRun.Status.FEASIBLE
        )
        ScheduledClass.objects.create(
            run=self.run,
            assignment=self.assignment,
            room=self.room,
            day=0,
            start_time="08:00",
            duration_slots=2,
            mode="sync",
        )

        self.url = f"/api/exports/schedules/{self.run.pk}.xlsx"

    def auth(self, token):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

    def _read_xlsx(self, response):
        """Consume a FileResponse's streaming content into an openpyxl workbook."""
        content = b"".join(response.streaming_content)
        return load_workbook(BytesIO(content))


class SchedulePdfExportTests(ExportTestCase):
    def setUp(self):
        super().setUp()
        self.pdf_url = f"/api/exports/schedules/{self.run.pk}.pdf"

    def test_unauthenticated_is_rejected(self):
        response = self.client.get(self.pdf_url)
        self.assertEqual(response.status_code, 401)

    def test_professor_is_forbidden(self):
        self.auth(self.prof_token)
        response = self.client.get(self.pdf_url)
        self.assertEqual(response.status_code, 403)

    def test_registrar_gets_pdf(self):
        self.auth(self.reg_token)
        response = self.client.get(self.pdf_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/pdf",
        )
        self.assertIn("attachment", response["Content-Disposition"])

    def test_unknown_run_returns_404(self):
        self.auth(self.reg_token)
        response = self.client.get("/api/exports/schedules/999999.pdf")
        self.assertEqual(response.status_code, 404)

    def test_pdf_has_valid_magic_bytes(self):
        self.auth(self.reg_token)
        response = self.client.get(self.pdf_url)
        content = b"".join(response.streaming_content)
        self.assertTrue(content.startswith(b"%PDF-"))
        self.assertTrue(content.rstrip().endswith(b"%%EOF"))

    def test_pdf_is_not_empty(self):
        self.auth(self.reg_token)
        response = self.client.get(self.pdf_url)
        content = b"".join(response.streaming_content)
        self.assertGreater(len(content), 1000)


class ExportTestCase(APITestCase):
    def setUp(self):
        # Registrar
        self.registrar = User.objects.create_user(
            username="reg", password="pass12345", is_staff=True
        )
        self.reg_token = Token.objects.create(user=self.registrar)

        # Professor
        self.dept = Department.objects.create(name="CS")
        self.prof_user = User.objects.create_user(
            username="prof1", password="pass12345"
        )
        self.prof = Professors.objects.create(
            user=self.prof_user, department=self.dept
        )
        self.prof_token = Token.objects.create(user=self.prof_user)

        # Schedule data
        self.subject = Subject.objects.create(
            code="CS101", title="Intro", units=3
        )
        self.section = Section.objects.create(name="BSIT-3A", headcount=30)
        self.room = Room.objects.create(name="R101", capacity=40)
        self.assignment = Assignment.objects.create(
            prof=self.prof, subject=self.subject, section=self.section
        )
        self.run = ScheduleRun.objects.create(
            algorithm="greedy", status=ScheduleRun.Status.FEASIBLE
        )
        ScheduledClass.objects.create(
            run=self.run,
            assignment=self.assignment,
            room=self.room,
            day=0,
            start_time="08:00",
            duration_slots=2,
            mode="sync",
        )

        self.url = f"/api/exports/schedules/{self.run.pk}.xlsx"

    def auth(self, token):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

    def _read_xlsx(self, response):
        """Consume a FileResponse's streaming content into an openpyxl workbook."""
        content = b"".join(response.streaming_content)
        return load_workbook(BytesIO(content))


class ScheduleExportTests(ExportTestCase):
    def test_unauthenticated_is_rejected(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 401)

    def test_professor_is_forbidden(self):
        self.auth(self.prof_token)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 403)

    def test_registrar_gets_xlsx(self):
        self.auth(self.reg_token)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", response["Content-Disposition"])

    def test_unknown_run_returns_404(self):
        self.auth(self.reg_token)
        response = self.client.get("/api/exports/schedules/999999.xlsx")
        self.assertEqual(response.status_code, 404)

    def test_xlsx_contains_expected_headers(self):
        self.auth(self.reg_token)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)

        wb = self._read_xlsx(response)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]

        self.assertEqual(
            headers,
            [
                "Day",
                "Start",
                "End",
                "Subject Code",
                "Title",
                "Section",
                "Professor",
                "Room",
                "Mode",
            ],
        )

    def test_xlsx_contains_schedule_row(self):
        self.auth(self.reg_token)
        response = self.client.get(self.url)

        wb = self._read_xlsx(response)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        # Header + 1 data row
        self.assertEqual(len(rows), 2)
        data = rows[1]
        self.assertEqual(data[0], "Monday")   # day=0
        self.assertEqual(data[1], "08:00")
        self.assertEqual(data[2], "10:00")    # 2 slots = 2 hours
        self.assertEqual(data[3], "CS101")
        self.assertEqual(data[6], "prof1")    # username fallback
        self.assertEqual(data[7], "R101")
        self.assertEqual(data[8], "Synchronous")
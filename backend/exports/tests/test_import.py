import io
import csv
import zipfile
from io import BytesIO
from datetime import time as dtime

from django.contrib.auth import get_user_model
from django.core.cache import cache
from openpyxl import load_workbook
from rest_framework.authtoken.models import Token
from rest_framework.test import APITestCase

from catalog.models import Department, Room, Section, Subject
from timetable.models import Assignment, MeetingSlot, AvailabilityWindow, BusyBlock
from users.models import Professors

from exports.import_utils import processor, parse_time, parse_day, parse_mode

User = get_user_model()


class ImportTestCase(APITestCase):
    def setUp(self):
        self.registrar = User.objects.create_user(
            username="reg", password="pass12345", is_staff=True
        )
        self.reg_token = Token.objects.create(user=self.registrar)

        self.prof_user = User.objects.create_user(username="prof1", password="pass12345")
        self.prof_user2 = User.objects.create_user(username="prof2", password="pass12345")
        self.dept = Department.objects.create(name="CS")
        self.prof = Professors.objects.create(user=self.prof_user, department=self.dept)
        self.prof2 = Professors.objects.create(user=self.prof_user2, department=self.dept)

        self.subject = Subject.objects.create(code="CS101", title="Intro", units=3)
        self.section = Section.objects.create(name="BSIT-3A", headcount=30)
        self.room = Room.objects.create(name="R101", capacity=40, department=self.dept)
        self.assignment = Assignment.objects.create(prof=self.prof, subject=self.subject, section=self.section)

        self.url_preview = "/api/imports/preview/"
        self.url_commit = "/api/imports/commit/"
        self.url_template = "/api/imports/template/"
        self.url_template_csv = "/api/imports/template/csv/"

    def auth(self, token):
        self.client.credentials(HTTP_AUTHORIZATION=f"Token {token.key}")

    def _make_xlsx(self, sheets_data: dict) -> BytesIO:
        """Create an XLSX file in memory from sheets data."""
        from openpyxl import Workbook
        wb = Workbook()
        first = True
        for sheet_name, rows in sheets_data.items():
            if first:
                ws = wb.active
                ws.title = sheet_name
                first = False
            else:
                ws = wb.create_sheet(sheet_name)
            for row_idx, row in enumerate(rows, start=1):
                for col_idx, val in enumerate(row, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=val)
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "test.xlsx"
        return buf

    def _make_csv_zip(self, sheets_data: dict) -> BytesIO:
        """Create a CSV ZIP file in memory from sheets data."""
        buf = BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for sheet_name, rows in sheets_data.items():
                output = io.StringIO()
                writer = csv.writer(output)
                for row in rows:
                    writer.writerow(row)
                zf.writestr(f"{sheet_name}.csv", output.getvalue())
        buf.seek(0)
        return buf


class ImportPreviewTests(ImportTestCase):
    def test_unauthenticated_rejected(self):
        response = self.client.post(self.url_preview, {"file": b"test"}, format="multipart")
        self.assertEqual(response.status_code, 401)

    def test_professor_forbidden(self):
        self.auth(Token.objects.create(user=self.prof_user))
        response = self.client.post(self.url_preview, {"file": b"test"}, format="multipart")
        self.assertEqual(response.status_code, 403)

    def test_no_file_rejected(self):
        self.auth(self.reg_token)
        response = self.client.post(self.url_preview, {}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertIn("No file provided", response.data["detail"])

    def test_invalid_format_rejected(self):
        self.auth(self.reg_token)
        response = self.client.post(self.url_preview, {"file": b"not a valid file"}, format="multipart")
        self.assertEqual(response.status_code, 400)

    def test_valid_xlsx_returns_preview(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [
                ["id", "name"],
                [self.dept.id, "Computer Science"],
            ],
            "Rooms": [
                ["id", "name", "capacity", "department_id"],
                [None, "M101", 30, self.dept.id],
            ],
            "Subjects": [
                ["id", "code", "title", "units"],
                [None, "MATH101", "Calculus", 3],
            ],
            "Sections": [
                ["id", "name", "headcount"],
                [None, "MATH-1A", 25],
            ],
            "Professors": [
                ["id", "user_id", "department_id", "max_daily_hours", "max_consecutive"],
                [None, self.prof_user.id, self.dept.id, 8, 4],
            ],
            "Assignments": [
                ["id", "professor_id", "subject_id", "section_id"],
            ],
            "Meetings": [
                ["id", "assignment_id", "order", "mode", "duration_slots"],
            ],
            "Availability": [
                ["id", "professor_id", "day", "start_time", "end_time", "is_preferred"],
            ],
            "BusyBlocks": [
                ["id", "professor_id", "day", "start_time", "end_time"],
            ],
        }
        xlsx = self._make_xlsx(sheets)
        response = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["valid"])
        self.assertIn("preview_token", response.data)
        self.assertEqual(response.data["summary"]["creates"], 4)
        self.assertEqual(response.data["summary"]["updates"], 1)

    def test_update_existing_records(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [
                ["id", "name"],
                [self.dept.id, "Computer Science"],
            ],
            "Rooms": [
                ["id", "name", "capacity", "department_id"],
                [self.room.id, "R101 Updated", 50, self.dept.id],
            ],
            "Subjects": [
                ["id", "code", "title", "units"],
                [self.subject.id, "CS101", "Introduction to CS", 4],
            ],
            "Sections": [
                ["id", "name", "headcount"],
                [self.section.id, "BSIT-3A", 35],
            ],
            "Professors": [
                ["id", "user_id", "department_id", "max_daily_hours", "max_consecutive"],
                [self.prof.id, self.prof_user.id, self.dept.id, 6, 3],
            ],
            "Assignments": [
                ["id", "professor_id", "subject_id", "section_id"],
            ],
            "Meetings": [
                ["id", "assignment_id", "order", "mode", "duration_slots"],
            ],
            "Availability": [
                ["id", "professor_id", "day", "start_time", "end_time", "is_preferred"],
            ],
            "BusyBlocks": [
                ["id", "professor_id", "day", "start_time", "end_time"],
            ],
        }
        xlsx = self._make_xlsx(sheets)
        response = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["valid"])
        self.assertEqual(response.data["summary"]["updates"], 5)
        self.assertEqual(response.data["summary"]["creates"], 0)

    def test_missing_required_field_error(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [
                ["id", "name"],
                [1, ""],  # missing name
            ],
            "Rooms": [["id", "name", "capacity"]],
            "Subjects": [["id", "code", "title"]],
            "Sections": [["id", "name", "headcount"]],
            "Professors": [["id", "user_id"]],
            "Assignments": [["id", "professor_id", "subject_id", "section_id"]],
            "Meetings": [["id", "assignment_id", "order", "mode"]],
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        xlsx = self._make_xlsx(sheets)
        response = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.data["valid"])
        self.assertTrue(any(e["field"] == "name" for e in response.data["errors"]))

    def test_invalid_mode_error(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [["id", "name"], [1, "CS"]],
            "Rooms": [["id", "name", "capacity"], [1, "R101", 40]],
            "Subjects": [["id", "code", "title"], [1, "CS101", "Intro"]],
            "Sections": [["id", "name", "headcount"], [1, "BSIT-3A", 30]],
            "Professors": [["id", "user_id"], [1, self.prof_user.id]],
            "Assignments": [["id", "professor_id", "subject_id", "section_id"], [1, 1, 1, 1]],
            "Meetings": [["id", "assignment_id", "order", "mode"], [1, 1, 1, "lecture"]],  # invalid mode
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        xlsx = self._make_xlsx(sheets)
        response = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.data["valid"])
        self.assertTrue(any(e["field"] == "mode" for e in response.data["errors"]))

    def test_missing_sync_meeting_error(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [["id", "name"], [1, "CS"]],
            "Rooms": [["id", "name", "capacity"], [1, "R101", 40]],
            "Subjects": [["id", "code", "title"], [1, "CS101", "Intro"]],
            "Sections": [["id", "name", "headcount"], [1, "BSIT-3A", 30]],
            "Professors": [["id", "user_id"], [1, self.prof_user.id]],
            "Assignments": [["id", "professor_id", "subject_id", "section_id"], [1, 1, 1, 1]],
            "Meetings": [["id", "assignment_id", "order", "mode"], [1, 1, 1, "async"]],  # only async
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        xlsx = self._make_xlsx(sheets)
        response = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.data["valid"])
        self.assertTrue(any("synchronous" in e["message"] for e in response.data["errors"]))

    def test_foreign_key_validation(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [["id", "name"], [1, "CS"]],
            "Rooms": [["id", "name", "capacity", "department_id"], [1, "R101", 40, 999]],  # invalid dept
            "Subjects": [["id", "code", "title"], [1, "CS101", "Intro"]],
            "Sections": [["id", "name", "headcount"], [1, "BSIT-3A", 30]],
            "Professors": [["id", "user_id"], [1, self.prof_user.id]],
            "Assignments": [["id", "professor_id", "subject_id", "section_id"], [1, 1, 1, 1]],
            "Meetings": [["id", "assignment_id", "order", "mode"], [1, 1, 1, "sync"]],
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        xlsx = self._make_xlsx(sheets)
        response = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.data["valid"])
        self.assertTrue(any(e["field"] == "department_id" for e in response.data["errors"]))

    def test_invalid_user_id_for_professor(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [["id", "name"], [1, "CS"]],
            "Rooms": [["id", "name", "capacity"], [1, "R101", 40]],
            "Subjects": [["id", "code", "title"], [1, "CS101", "Intro"]],
            "Sections": [["id", "name", "headcount"], [1, "BSIT-3A", 30]],
            "Professors": [["id", "user_id"], [1, 99999]],  # non-existent user
            "Assignments": [["id", "professor_id", "subject_id", "section_id"]],
            "Meetings": [["id", "assignment_id", "order", "mode"]],
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        xlsx = self._make_xlsx(sheets)
        response = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.data["valid"])


class ImportCommitTests(ImportTestCase):
    def test_commit_creates_records(self):
        self.auth(self.reg_token)
        # Create a new user without a professor profile
        new_user = User.objects.create_user(username="newprof", password="pass12345")
        sheets = {
            "Departments": [["id", "name"], [None, "Math"]],
            "Rooms": [["id", "name", "capacity"], [None, "M101", 30]],
            "Subjects": [["id", "code", "title"], [None, "MATH101", "Calculus"]],
            "Sections": [["id", "name", "headcount"], [None, "MATH-1A", 25]],
            "Professors": [["id", "user_id"], [None, new_user.id]],
            "Assignments": [["id", "professor_id", "subject_id", "section_id"]],
            "Meetings": [["id", "assignment_id", "order", "mode"]],
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        xlsx = self._make_xlsx(sheets)
        preview_resp = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertTrue(preview_resp.data["valid"])
        token = preview_resp.data["preview_token"]

        commit_resp = self.client.post(self.url_commit, {"preview_token": token}, format="json")
        
        self.assertEqual(commit_resp.status_code, 200)
        self.assertTrue(commit_resp.data["committed"])
        self.assertEqual(commit_resp.data["summary"]["creates"], 5)

        self.assertTrue(Department.objects.filter(name="Math").exists())
        self.assertTrue(Room.objects.filter(name="M101").exists())
        self.assertTrue(Subject.objects.filter(code="MATH101").exists())
        self.assertTrue(Section.objects.filter(name="MATH-1A").exists())
        self.assertTrue(Professors.objects.filter(user=new_user).exists())

    def test_commit_updates_records(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [["id", "name"], [self.dept.id, "Computer Science"]],
            "Rooms": [["id", "name", "capacity"], [self.room.id, "R101 Updated", 50]],
            "Subjects": [["id", "code", "title"], [self.subject.id, "CS101", "Intro to CS"]],
            "Sections": [["id", "name", "headcount"], [self.section.id, "BSIT-3A", 35]],
            "Professors": [["id", "user_id"], [self.prof.id, self.prof_user.id]],
            "Assignments": [["id", "professor_id", "subject_id", "section_id"]],
            "Meetings": [["id", "assignment_id", "order", "mode"]],
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        xlsx = self._make_xlsx(sheets)
        preview_resp = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertTrue(preview_resp.data["valid"])
        token = preview_resp.data["preview_token"]

        commit_resp = self.client.post(self.url_commit, {"preview_token": token}, format="json")
        self.assertEqual(commit_resp.status_code, 200)
        self.assertEqual(commit_resp.data["summary"]["updates"], 5)

        self.dept.refresh_from_db()
        self.assertEqual(self.dept.name, "Computer Science")
        self.room.refresh_from_db()
        self.assertEqual(self.room.name, "R101 Updated")
        self.assertEqual(self.room.capacity, 50)

    def test_commit_expired_token_rejected(self):
        self.auth(self.reg_token)
        response = self.client.post(self.url_commit, {"preview_token": "invalid-token"}, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.data["committed"])

    def test_commit_missing_token_rejected(self):
        self.auth(self.reg_token)
        response = self.client.post(self.url_commit, {}, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("preview_token is required", response.data["detail"])

    def test_commit_atomic_rollback_on_error(self):
        self.auth(self.reg_token)
        sheets = {
            "Departments": [["id", "name"], [1, "Valid Dept"]],
            "Rooms": [["id", "name", "capacity"], [1, "R101", 40]],
            "Subjects": [["id", "code", "title"], [1, "CS101", "Intro"]],
            "Sections": [["id", "name", "headcount"], [1, "BSIT-3A", 30]],
            "Professors": [["id", "user_id"], [1, self.prof_user.id]],
            "Assignments": [["id", "professor_id", "subject_id", "section_id"], [1, 1, 1, 1]],
            "Meetings": [["id", "assignment_id", "order", "mode"], [1, 1, 1, "sync"]],
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        xlsx = self._make_xlsx(sheets)
        preview_resp = self.client.post(self.url_preview, {"file": xlsx}, format="multipart")
        self.assertTrue(preview_resp.data["valid"])
        token = preview_resp.data["preview_token"]

        # Delete the professor user to cause commit failure
        self.prof_user.delete()

        commit_resp = self.client.post(self.url_commit, {"preview_token": token}, format="json")
        self.assertEqual(commit_resp.status_code, 400)
        self.assertFalse(commit_resp.data["committed"])

        # Verify nothing was created (rollback)
        self.assertFalse(Department.objects.filter(name="Valid Dept").exists())


class ImportTemplateTests(ImportTestCase):
    def test_xlsx_template_download(self):
        self.auth(self.reg_token)
        response = self.client.get(self.url_template)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("schedule_import_template.xlsx", response["Content-Disposition"])

        content = b"".join(response.streaming_content)
        wb = load_workbook(BytesIO(content))
        self.assertEqual(set(wb.sheetnames), set([
            "Departments", "Rooms", "Subjects", "Sections", "Professors",
            "Assignments", "Meetings", "Availability", "BusyBlocks", "Instructions"
        ]))

        for sheet_name in ["Departments", "Rooms", "Subjects", "Sections", "Professors",
                           "Assignments", "Meetings", "Availability", "BusyBlocks"]:
            ws = wb[sheet_name]
            self.assertEqual(ws.max_row, 2)  # header + example row

    def test_csv_template_download(self):
        self.auth(self.reg_token)
        response = self.client.get(self.url_template_csv)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        self.assertIn("schedule_import_template.zip", response["Content-Disposition"])

        content = b"".join(response.streaming_content)
        with zipfile.ZipFile(BytesIO(content), "r") as zf:
            names = zf.namelist()
            for sheet in ["Departments", "Rooms", "Subjects", "Sections", "Professors",
                          "Assignments", "Meetings", "Availability", "BusyBlocks"]:
                self.assertIn(f"{sheet}.csv", names)
            self.assertIn("Instructions.txt", names)

    def test_professor_cannot_access_template(self):
        self.auth(Token.objects.create(user=self.prof_user))
        response = self.client.get(self.url_template)
        self.assertEqual(response.status_code, 403)


class ImportProcessorUnitTests(ImportTestCase):
    def test_parse_time_formats(self):
        self.assertEqual(parse_time("08:00"), dtime(8, 0))
        self.assertEqual(parse_time("13:30"), dtime(13, 30))
        self.assertIsNone(parse_time(""))
        self.assertIsNone(parse_time(None))

    def test_parse_day_formats(self):
        self.assertEqual(parse_day(0), 0)
        self.assertEqual(parse_day("Monday"), 0)
        self.assertEqual(parse_day("Mon"), 0)
        self.assertEqual(parse_day("5"), 5)
        self.assertIsNone(parse_day(""))
        self.assertIsNone(parse_day(None))
        self.assertIsNone(parse_day("Invalid"))

    def test_parse_mode(self):
        self.assertEqual(parse_mode("async"), "async")
        self.assertEqual(parse_mode("SYNC"), "sync")
        self.assertEqual(parse_mode("Lab"), "lab")
        self.assertIsNone(parse_mode("lecture"))
        self.assertIsNone(parse_mode(""))

    def test_generate_template_xlsx(self):
        content = processor.generate_template_xlsx()
        wb = load_workbook(BytesIO(content))
        self.assertEqual(set(wb.sheetnames), set([
            "Departments", "Rooms", "Subjects", "Sections", "Professors",
            "Assignments", "Meetings", "Availability", "BusyBlocks", "Instructions"
        ]))

    def test_generate_template_csv_zip(self):
        content = processor.generate_template_csv_zip()
        with zipfile.ZipFile(BytesIO(content), "r") as zf:
            names = zf.namelist()
            for sheet in ["Departments", "Rooms", "Subjects", "Sections", "Professors",
                          "Assignments", "Meetings", "Availability", "BusyBlocks"]:
                self.assertIn(f"{sheet}.csv", names)
            self.assertIn("Instructions.txt", names)

    def test_csv_zip_parsing(self):
        sheets = {
            "Departments": [["id", "name"], ["", "Physics"]],
            "Rooms": [["id", "name", "capacity"], ["", "P101", "25"]],
            "Subjects": [["id", "code", "title"], ["", "PHYS101", "Physics"]],
            "Sections": [["id", "name", "headcount"], ["", "PHYS-1A", "20"]],
            "Professors": [["id", "user_id"], ["", str(self.prof_user2.id)]],
            "Assignments": [["id", "professor_id", "subject_id", "section_id"]],
            "Meetings": [["id", "assignment_id", "order", "mode"]],
            "Availability": [["id", "professor_id", "day", "start_time", "end_time"]],
            "BusyBlocks": [["id", "professor_id", "day", "start_time", "end_time"]],
        }
        csv_zip = self._make_csv_zip(sheets)
        csv_zip.name = "test.zip"
        parsed = processor.parse_file(csv_zip)
        self.assertEqual(len(parsed["Departments"]), 1)
        self.assertEqual(parsed["Departments"][0]["name"], "Physics")